from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\Projects\tool-craw")
WORK = ROOT / "work"
OUT_DIR = ROOT / "output" / "docx"
ASSET_DIR = WORK / "proposal_v2_assets"
OUT_FILE = OUT_DIR / "Ban_de_xuat_cai_tien_he_thong_du_bao_ngap_V2.docx"

WORK.mkdir(parents=True, exist_ok=True)
OUT_DIR.mkdir(parents=True, exist_ok=True)
ASSET_DIR.mkdir(parents=True, exist_ok=True)

NAVY = "17365D"
BLUE = "2F75B5"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F4F6F7"
LIGHT_BORDER = "D9D9D9"
BLACK = "000000"
WHITE = "FFFFFF"


FONT_REG = r"C:\Windows\Fonts\arial.ttf"
FONT_BOLD = r"C:\Windows\Fonts\arialbd.ttf"


def font(size, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REG, size)


def centered_multiline(draw, box, text, fnt, fill, spacing=6):
    x1, y1, x2, y2 = box
    bbox = draw.multiline_textbbox((0, 0), text, font=fnt, spacing=spacing, align="center")
    tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
    draw.multiline_text(((x1 + x2 - tw) / 2, (y1 + y2 - th) / 2), text,
                        font=fnt, fill=fill, spacing=spacing, align="center")


def draw_box(draw, box, title, body, color="#2F75B5"):
    x1, y1, x2, y2 = box
    draw.rounded_rectangle(box, radius=18, fill="#F7FAFC", outline=color, width=3)
    title_box = (x1 + 12, y1 + 12, x2 - 12, y1 + int((y2-y1)*0.48))
    body_box = (x1 + 12, y1 + int((y2-y1)*0.43), x2 - 12, y2 - 12)
    centered_multiline(draw, title_box, title, font(24, True), "#17365D", spacing=4)
    centered_multiline(draw, body_box, body, font(18), "#202124", spacing=5)


def arrow(draw, start, end, color="#5B6573", dashed=False):
    x1, y1 = start
    x2, y2 = end
    if dashed:
        steps = 12
        for i in range(steps):
            if i % 2 == 0:
                a = i / steps
                b = (i + 1) / steps
                draw.line((x1 + (x2-x1)*a, y1 + (y2-y1)*a,
                           x1 + (x2-x1)*b, y1 + (y2-y1)*b), fill=color, width=3)
    else:
        draw.line((x1, y1, x2, y2), fill=color, width=3)
    import math
    angle = math.atan2(y2-y1, x2-x1)
    length = 15
    p1 = (x2 - length * math.cos(angle - 0.55), y2 - length * math.sin(angle - 0.55))
    p2 = (x2 - length * math.cos(angle + 0.55), y2 - length * math.sin(angle + 0.55))
    draw.polygon([(x2, y2), p1, p2], fill=color)


def make_architecture():
    img = Image.new("RGB", (2400, 1280), "white")
    draw = ImageDraw.Draw(img)
    draw.text((50, 45), "Kiến trúc hệ thống dự báo ngập cải tiến", font=font(40, True), fill="#111111")
    draw.text((50, 100), "Tách dự báo ngập cục bộ 15 đến 30 phút khỏi cảnh báo bối cảnh 24 giờ",
              font=font(25), fill="#4B5563")
    boxes = {
        "sensor": (70, 270, 480, 530),
        "external": (70, 700, 480, 960),
        "esp": (650, 270, 1060, 530),
        "store": (1230, 270, 1640, 530),
        "local": (1810, 270, 2320, 530),
        "context": (1230, 700, 1640, 960),
        "dash": (1810, 700, 2320, 960),
    }
    draw_box(draw, boxes["sensor"], "Cảm biến tại trường", "Mưa RS485\nMực nước\nChu kỳ 1 đến 3 phút")
    draw_box(draw, boxes["external"], "Dữ liệu bên ngoài", "Dự báo mưa\nMực nước hoặc lưu lượng sông\nThủy triều khi có")
    draw_box(draw, boxes["esp"], "ESP32 và truyền dữ liệu", "Đọc Modbus RTU\nKiểm tra hợp lệ\nHTTP hoặc MQTT")
    draw_box(draw, boxes["store"], "Kho dữ liệu chuỗi thời gian", "Dữ liệu thô bất biến\nDữ liệu đã làm sạch\nCờ chất lượng")
    draw_box(draw, boxes["local"], "Mô hình ngập cục bộ", "Dự báo sau 15 đến 30 phút\nƯu tiên Recall lớp nguy hiểm")
    draw_box(draw, boxes["context"], "Mô hình bối cảnh", "Nguy cơ 6 đến 24 giờ\nChỉ hỗ trợ quyết định")
    draw_box(draw, boxes["dash"], "Dashboard và cảnh báo", "Xanh Vàng Đỏ\nTelegram\nNhật ký và xác nhận")
    arrow(draw, (480, 400), (650, 400))
    arrow(draw, (1060, 400), (1230, 400))
    arrow(draw, (1640, 400), (1810, 400))
    arrow(draw, (480, 830), (1230, 830), dashed=True)
    arrow(draw, (1435, 530), (1435, 700), dashed=True)
    arrow(draw, (1640, 830), (1810, 830), dashed=True)
    arrow(draw, (2065, 530), (2065, 700))
    centered_multiline(draw, (200, 1050, 2200, 1190),
                       "Nguyên tắc vận hành  Dữ liệu cảm biến quyết định cảnh báo gần\nDữ liệu bên ngoài bổ sung bối cảnh",
                       font(23), "#374151", spacing=7)
    path = ASSET_DIR / "kien_truc_v2.png"
    img.save(path, dpi=(220, 220))
    return path


def make_experiment_flow():
    img = Image.new("RGB", (2400, 940), "white")
    draw = ImageDraw.Draw(img)
    draw.text((50, 42), "Quy trình nghiên cứu và kiểm định", font=font(40, True), fill="#111111")
    items = [
        ("Khảo sát", "Chọn điểm ngập\nXác định ngưỡng thực địa"),
        ("Hiệu chuẩn", "Kiểm tra mưa\nKiểm tra mực nước"),
        ("Thu thập", "Dữ liệu 1 đến 3 phút\nGhi cờ chất lượng"),
        ("Gán nhãn", "Mực nước tương lai\n15 và 30 phút"),
        ("Huấn luyện", "Baseline và ba mô hình\nKhông trộn tương lai"),
        ("Kiểm định", "Tách theo sự kiện\nĐánh giá hiện trường"),
    ]
    xs = [50, 445, 840, 1235, 1630, 2025]
    for i, ((title, body), x) in enumerate(zip(items, xs)):
        box = (x, 280, x + 315, 610)
        draw_box(draw, box, title, body, color="#2F75B5")
        centered_multiline(draw, (x, 190, x + 315, 255), str(i + 1), font(25, True), "#2F75B5")
        if i < len(items) - 1:
            arrow(draw, (x + 315, 445), (xs[i + 1], 445))
    centered_multiline(draw, (150, 700, 2250, 850),
                       "Dữ liệu giả lập chỉ dùng để kiểm tra kỹ thuật hoặc hỗ trợ huấn luyện\nKhông dùng làm tập kiểm định độc lập",
                       font(23), "#7A2E2E", spacing=7)
    path = ASSET_DIR / "quy_trinh_nghien_cuu.png"
    img.save(path, dpi=(220, 220))
    return path


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_table_header(row):
    repeat_header(row)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_run_font(run, name="Arial", size=10.5, bold=None, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(14)
    title_ppr = title.element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, before, after in (
        ("Heading 1", 16, 14, 7),
        ("Heading 2", 13, 11, 5),
        ("Heading 3", 11.5, 8, 4),
    ):
        st = styles[name]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True


def add_field(paragraph, instr):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instr
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color="666666")


def add_paragraph(doc, text="", bold_lead=None, align=None, keep=False):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    if align is not None:
        p.alignment = align
    p.paragraph_format.keep_together = keep
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    r = p.add_run(text)
    set_run_font(r)
    p.paragraph_format.space_after = Pt(3)
    return p


def add_number(doc, text, level=0):
    p = doc.add_paragraph(style="List Number" if level == 0 else "List Number 2")
    r = p.add_run(text)
    set_run_font(r)
    p.paragraph_format.space_after = Pt(3)
    return p


def add_table(doc, headers, rows, widths=None, header_fill=NAVY, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        c = hdr.cells[i]
        c.text = ""
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(c, header_fill)
        set_cell_margins(c)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        set_run_font(r, size=font_size, bold=True, color=WHITE)
        if widths:
            c.width = Inches(widths[i])
    for ridx, row in enumerate(rows):
        added_row = table.add_row()
        prevent_row_split(added_row)
        cells = added_row.cells
        shade = PALE_BLUE if ridx % 2 else WHITE
        for i, value in enumerate(row):
            c = cells[i]
            c.text = ""
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_shading(c, shade)
            set_cell_margins(c)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
            if widths:
                c.width = Inches(widths[i])
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_figure(doc, path, caption, width=6.8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    cap.paragraph_format.keep_together = True
    r = cap.add_run(caption)
    set_run_font(r, size=9, bold=False, color="555555")
    r.italic = True


def new_page(doc):
    doc.add_page_break()


arch_path = make_architecture()
flow_path = make_experiment_flow()

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.70)
section.bottom_margin = Inches(0.70)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)
configure_styles(doc)

# Footer with page number
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run("Trang ")
set_run_font(fr, size=9, color="666666")
add_field(fp, "PAGE")

# Cover page
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(20)
r = p.add_run("TRƯỜNG THPT TRẦN QUÝ CÁP")
set_run_font(r, size=12, bold=True)

for _ in range(2):
    doc.add_paragraph()

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("BẢN ĐỀ XUẤT CẢI TIẾN HỆ THỐNG DỰ BÁO NGẬP CỤC BỘ TẠI TRƯỜNG HỌC")
set_run_font(r, size=24, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Phiên bản 2")
set_run_font(r, size=16, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(10)
r = p.add_run("Hội thi khoa học kỹ thuật STEM cấp trường năm học 2026 2027")
set_run_font(r, size=12, bold=False)

for _ in range(3):
    doc.add_paragraph()

cover_rows = [
    ("Lĩnh vực", "Hệ thống nhúng  Trí tuệ nhân tạo  Môi trường"),
    ("Nhóm thực hiện", "Nguyễn Tấn Phúc và Ngô Nguyên Khôi  Lớp 10 23"),
    ("Đơn vị", "Trường THPT Trần Quý Cáp  Đà Nẵng"),
    ("Phạm vi", "Một điểm ngập cục bộ trong khuôn viên trường"),
]
t = doc.add_table(rows=0, cols=2)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
for left, right in cover_rows:
    cells = t.add_row().cells
    cells[0].width = Inches(1.35)
    cells[1].width = Inches(4.85)
    for c in cells:
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(c, top=80, bottom=80)
    p0 = cells[0].paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r0 = p0.add_run(left)
    set_run_font(r0, size=10.5, bold=True)
    p1 = cells[1].paragraphs[0]
    r1 = p1.add_run(right)
    set_run_font(r1, size=10.5)

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Tháng 10 năm 2026")
set_run_font(r, size=11)

new_page(doc)

doc.add_heading("Tóm tắt đề xuất cải tiến", level=1)
add_paragraph(doc,
    "Phiên bản 2 giữ mục tiêu của đề tài ban đầu là dự báo sớm ngập cục bộ tại trường học, "
    "nhưng thu hẹp bài toán chính về khoảng dự báo 15 đến 30 phút và tách riêng lớp cảnh báo "
    "bối cảnh 6 đến 24 giờ. Cách tổ chức này phù hợp hơn với thiết bị hiện có, khả năng thu thập "
    "dữ liệu của nhóm và yêu cầu kiểm chứng trong phạm vi một điểm đo tại trường.")
add_paragraph(doc,
    "Kết luận chính là hệ thống chỉ có thể dự báo ngập cục bộ đáng tin cậy khi bổ sung cảm biến "
    "mực nước và gán nhãn bằng mực nước tương lai thực đo. Dữ liệu mô phỏng hoặc mực nước do công "
    "thức sinh ra có thể hỗ trợ kiểm tra kỹ thuật, nhưng không được dùng làm bằng chứng đánh giá "
    "độc lập cho hiệu năng của mô hình.")

doc.add_heading("Những thay đổi trọng tâm", level=2)
changes = [
    "Tách dự báo ngập cục bộ 15 đến 30 phút khỏi cảnh báo thủy văn 6 đến 24 giờ.",
    "Định nghĩa nhãn bằng trạng thái ngập trong tương lai thay vì suy ra từ mực nước hiện tại.",
    "Bổ sung cảm biến mực nước và chuẩn hóa cửa sổ thời gian của các đặc trưng.",
    "Đánh giá theo thứ tự thời gian và theo từng sự kiện mưa, không xáo trộn ngẫu nhiên toàn bộ chuỗi.",
    "Dùng Recall lớp nguy hiểm, Precision, F1 Macro, thời gian báo trước và số cảnh báo giả làm chỉ số chính.",
    "Tách dữ liệu quan trắc, dữ liệu ước lượng và dữ liệu giả lập trong toàn bộ pipeline.",
]
for x in changes:
    add_bullet(doc, x)

doc.add_heading("1 Phân tích phiên bản 1", level=1)
doc.add_heading("1 1 Điểm phù hợp", level=2)
for x in [
    "Bài toán xuất phát từ nhu cầu thực tế tại trường và có đối tượng sử dụng rõ ràng.",
    "Kiến trúc cảm biến ESP32 dữ liệu mô hình và dashboard tạo thành một chuỗi sản phẩm có thể trình diễn.",
    "Phiên bản 1 đã lựa chọn đúng nhóm thuật toán bảng dễ giải thích gồm Logistic Regression Decision Tree và Random Forest.",
    "Việc ưu tiên Recall cho lớp nguy hiểm phù hợp với bài toán cảnh báo an toàn.",
]:
    add_bullet(doc, x)

doc.add_heading("1 2 Điểm cần điều chỉnh", level=2)
analysis_rows = [
    ("Phạm vi dự báo", "Tài liệu đặt mục tiêu 15 đến 30 phút nhưng phần mềm hiện tại huấn luyện nhãn 24 giờ. Hai bài toán dùng cửa sổ dữ liệu và cách đánh giá khác nhau."),
    ("Nguồn nhãn", "Dataset hiện có 366 ngày. Khoảng 52,46 phần trăm mực nước được sinh từ công thức thủy lực nội bộ; toàn bộ 20 mẫu lớp nguy hiểm cao thuộc nhóm này."),
    ("Độ tin cậy", "Recall 100 phần trăm hiện chỉ dựa trên rất ít mẫu lớp nguy hiểm và chưa phải tập sự kiện thực địa độc lập."),
    ("Đặc trưng thời gian", "Delta trong dữ liệu huấn luyện tính theo ngày, trong khi dashboard có thể tính giữa hai lần bấm chỉ cách nhau vài giây."),
    ("Phần cứng", "Nhóm đã có cảm biến mưa RS485 nhưng chưa có cảm biến mực nước nên chưa đo được biến quan trọng nhất của ngập cục bộ."),
    ("Kết nối", "Mô phỏng ESP32 và dashboard đang hoạt động độc lập; chưa có cầu nối HTTP MQTT hoặc Serial vận hành liên tục."),
]
add_table(doc, ["Nội dung", "Nhận định"], analysis_rows, widths=[1.45, 5.35], font_size=9.3)

doc.add_heading("1 3 Hướng cải tiến được chọn", level=2)
add_paragraph(doc,
    "Đề tài nên tập trung vào nowcasting ngập cục bộ tại một vị trí trong trường với thời gian báo trước "
    "15 và 30 phút. Cảnh báo 6 đến 24 giờ từ dự báo mưa và dữ liệu thủy văn chỉ đóng vai trò bối cảnh "
    "để nhà trường chuẩn bị, không thay thế kết quả cảm biến tại chỗ.")

doc.add_heading("2 Tên đề tài và phạm vi nghiên cứu", level=1)
add_paragraph(doc,
    "Tên đề tài đề xuất  Nghiên cứu và xây dựng hệ thống IoT kết hợp học máy dự báo nguy cơ ngập cục bộ "
    "trước 15 đến 30 phút tại một vị trí trong khuôn viên Trường THPT Trần Quý Cáp.")

doc.add_heading("2 1 Bài toán chính", level=2)
add_paragraph(doc,
    "Tại thời điểm t hệ thống sử dụng chuỗi mưa và mực nước gần nhất để dự báo trạng thái ngập tại thời "
    "điểm t cộng 15 phút và t cộng 30 phút. Kết quả gồm ba mức An toàn Cảnh báo và Nguy hiểm cao.")

doc.add_heading("2 2 Phạm vi giới hạn", level=2)
for x in [
    "Một vị trí đo đã khảo sát trong khuôn viên trường.",
    "Dự báo ngập mặt hoặc đầy rãnh thoát nước tại điểm đo, không tuyên bố dự báo lũ cho toàn thành phố.",
    "Prototype phục vụ nghiên cứu và diễn tập, không tự động ra quyết định sơ tán hoặc ngắt điện.",
    "Dữ liệu bên ngoài chỉ được dùng khi có dấu thời gian, nguồn và cờ chất lượng rõ ràng.",
]:
    add_bullet(doc, x)

doc.add_heading("3 Câu hỏi nghiên cứu và giả thuyết", level=1)
add_paragraph(doc,
    "Câu hỏi nghiên cứu  Chuỗi cường độ mưa lượng mưa tích lũy mực nước và tốc độ dâng nước có giúp "
    "mô hình dự báo đúng trạng thái ngập tại điểm đo trước 15 đến 30 phút trên các sự kiện mưa chưa từng "
    "được dùng để huấn luyện hay không.")
add_paragraph(doc,
    "Giả thuyết  Khi dữ liệu cảm biến được lấy mẫu ổn định và nhãn được xác định bằng mực nước tương lai "
    "thực đo mô hình đa biến sẽ nhận diện lớp nguy hiểm tốt hơn phương pháp chỉ dùng ngưỡng mưa hoặc "
    "ngưỡng mực nước hiện tại. Giả thuyết được chấp nhận hoặc bác bỏ dựa trên tập kiểm định theo thời gian.")

doc.add_heading("4 Mục tiêu nghiên cứu", level=1)
doc.add_heading("4 1 Mục tiêu tổng quát", level=2)
add_paragraph(doc,
    "Xây dựng và kiểm định một prototype IoT có khả năng đo mưa đo mực nước truyền dữ liệu không dây và "
    "dự báo nguy cơ ngập cục bộ trước 15 đến 30 phút tại một điểm đo trong trường.")
doc.add_heading("4 2 Mục tiêu cụ thể", level=2)
for x in [
    "Hoàn thiện trạm đo sử dụng ESP32 cảm biến mưa RS485 và cảm biến mực nước.",
    "Xây dựng dữ liệu chuỗi thời gian có cờ chất lượng và phân biệt dữ liệu thực với dữ liệu giả lập.",
    "So sánh baseline ngưỡng với Logistic Regression Decision Tree và Random Forest.",
    "Đánh giá mô hình trên các sự kiện mưa độc lập theo thứ tự thời gian.",
    "Hiển thị dữ liệu cảnh báo và tình trạng thiết bị trên dashboard kèm nhật ký Telegram.",
]:
    add_bullet(doc, x)

doc.add_heading("5 Kiến trúc kỹ thuật", level=1)
add_figure(doc, arch_path, "Hình 1  Kiến trúc hai lớp của hệ thống cải tiến", width=6.85)
add_paragraph(doc,
    "Lớp gần sử dụng dữ liệu cảm biến tại trường để dự báo 15 đến 30 phút. Lớp bối cảnh sử dụng dữ liệu "
    "khí tượng thủy văn bên ngoài để đưa ra xu thế 6 đến 24 giờ. Dashboard phải hiển thị rõ nguồn và "
    "khoảng dự báo để người dùng không nhầm hai kết quả.")

doc.add_heading("6 Input đầu vào và nhãn dự báo", level=1)
input_rows = [
    ("Mưa hiện tại", "Cảm biến mưa RS485", "1 đến 3 phút", "Cường độ mưa và lượng mưa từng khoảng"),
    ("Mưa tích lũy", "Tính từ lịch sử", "5 15 30 60 phút", "Phản ánh lượng nước đưa vào hệ thống thoát nước"),
    ("Mực nước", "Cảm biến tại điểm ngập", "1 đến 3 phút", "Biến trạng thái trực tiếp của ngập cục bộ"),
    ("Tốc độ dâng", "Tính từ mực nước", "5 và 15 phút", "Phát hiện nước tăng nhanh"),
    ("Độ ổn định cảm biến", "ESP32", "Mỗi bản ghi", "Mã lỗi mất kết nối hoặc giá trị ngoài dải"),
    ("Dự báo mưa", "Nguồn bên ngoài", "Theo giờ", "Input phụ cho lớp cảnh báo bối cảnh"),
    ("Lưu lượng hoặc mực sông", "Nguồn thủy văn", "Theo ngày hoặc giờ", "Input phụ khi có quan hệ với điểm trường"),
]
add_table(doc, ["Nhóm input", "Nguồn", "Cửa sổ", "Ý nghĩa"], input_rows,
          widths=[1.25, 1.45, 1.30, 2.80], font_size=8.8)

doc.add_heading("6 1 Định nghĩa nhãn", level=2)
add_paragraph(doc,
    "Nhãn tại thời điểm t phải được xác định từ mực nước thực đo ở tương lai t cộng h. Với h bằng 15 hoặc "
    "30 phút hệ thống phân loại theo hai ngưỡng thực địa do nhóm khảo sát tại điểm đặt cảm biến. Ngưỡng "
    "của trạm thủy văn Câu Lâu không được dùng thay cho ngưỡng ngập trong sân trường.")
label_rows = [
    ("0 An toàn", "Mực nước tương lai dưới ngưỡng cảnh báo tại điểm đo"),
    ("1 Cảnh báo", "Mực nước tương lai đạt vùng cần theo dõi hoặc đang tăng nhanh"),
    ("2 Nguy hiểm cao", "Mực nước tương lai vượt ngưỡng nguy hiểm đã khảo sát"),
]
add_table(doc, ["Nhãn", "Quy tắc thực địa"], label_rows, widths=[1.55, 5.25], font_size=9.3)

doc.add_heading("6 2 Nguyên tắc chống rò rỉ dữ liệu", level=2)
for x in [
    "Chỉ dùng thông tin có sẵn tại hoặc trước thời điểm t làm input.",
    "Không dùng mực nước tương lai trong feature engineering.",
    "Dòng cuối không có đủ tương lai phải loại khỏi tập huấn luyện.",
    "Chuẩn hóa và điền thiếu phải được học trên tập train rồi áp dụng cho validation và test.",
    "Một sự kiện mưa không được chia đồng thời vào train và test.",
]:
    add_bullet(doc, x)

doc.add_heading("7 Thiết bị phần cứng", level=1)
doc.add_heading("7 1 Thiết bị hiện có", level=2)
hardware_rows = [
    ("ESP32 NodeMCU WiFi", "Đủ", "Điều khiển trung tâm và truyền dữ liệu"),
    ("Mạch mở rộng ESP32 có nguồn DC", "Cần xác minh", "Kiểm tra dải điện áp DC trước khi cấp 12V"),
    ("Cảm biến mưa RS485", "Đủ cho đo mưa", "Cần datasheet Modbus địa chỉ thanh ghi và hệ số quy đổi"),
    ("Module RS485 TTL", "Cần xác minh", "Ưu tiên mức logic 3,3V hoặc dùng chuyển mức"),
    ("Adapter 12V", "Có thể dùng", "Chỉ cấp vào phần cứng hỗ trợ 12V"),
    ("Dây cái cái", "Đủ thử bàn", "Không phù hợp lắp ngoài trời lâu dài"),
]
add_table(doc, ["Thiết bị", "Đánh giá", "Lưu ý"], hardware_rows,
          widths=[2.05, 1.25, 3.50], font_size=9.0)

doc.add_heading("7 2 Thiết bị cần bổ sung", level=2)
for x in [
    "Cảm biến mực nước chống nước phù hợp khoảng đo thực tế.",
    "Bộ hạ áp 12V xuống 5V nếu mạch mở rộng không nhận trực tiếp 12V.",
    "Module RS485 3,3V hoặc bộ chuyển mức logic khi module hiện tại xuất TTL 5V.",
    "Hộp kỹ thuật chống nước đầu nối chắc chắn cầu chì và bảo vệ quá áp.",
    "Dây xoắn đôi cho đường A B và điện trở kết thúc khi cáp dài.",
]:
    add_bullet(doc, x)

doc.add_heading("8 Thiết kế thu thập dữ liệu", level=1)
add_figure(doc, flow_path, "Hình 2  Quy trình nghiên cứu và kiểm định dữ liệu", width=6.85)

doc.add_heading("8 1 Chu kỳ và cấu trúc bản ghi", level=2)
record_rows = [
    ("timestamp", "Thời gian ISO có múi giờ", "Bắt buộc"),
    ("device_id", "Mã trạm đo", "Bắt buộc"),
    ("rain_mm", "Mưa trong chu kỳ", "Thực đo"),
    ("rain_rate_mmh", "Cường độ mưa", "Thực đo hoặc tính theo datasheet"),
    ("water_level_cm", "Mực nước tại điểm đo", "Thực đo"),
    ("signal_quality", "Trạng thái WiFi RS485 và lỗi cảm biến", "Cờ chất lượng"),
    ("source", "observed simulated external estimated", "Bắt buộc"),
]
add_table(doc, ["Trường", "Nội dung", "Yêu cầu"], record_rows,
          widths=[1.55, 3.65, 1.60], font_size=9.0)

doc.add_heading("8 2 Kế hoạch tạo dữ liệu", level=2)
for x in [
    "Hiệu chuẩn cảm biến bằng phép đo đối chứng trước khi lắp ngoài trời.",
    "Ghi dữ liệu liên tục theo chu kỳ 1 đến 3 phút trong cả ngày khô và ngày mưa.",
    "Ghi nhật ký các sự kiện như mất điện vệ sinh cảm biến nghẽn rãnh hoặc thay vị trí thiết bị.",
    "Tạo thí nghiệm bể nước hoặc máng thoát nước để kiểm tra cảm biến và pipeline nhưng gắn nhãn simulated.",
    "Chỉ sử dụng sự kiện thực địa chưa xuất hiện trong train để báo cáo hiệu năng cuối cùng.",
]:
    add_number(doc, x)

doc.add_heading("9 Mô hình và phương pháp đánh giá", level=1)
model_rows = [
    ("Baseline ngưỡng", "Mưa mực nước và tốc độ dâng", "Mốc so sánh dễ giải thích"),
    ("Logistic Regression", "Feature đã chuẩn hóa", "Xác suất và hệ số dễ giải thích"),
    ("Decision Tree", "Giới hạn độ sâu", "Luật quyết định trực quan"),
    ("Random Forest", "Class weight và giới hạn độ sâu", "Nắm bắt quan hệ phi tuyến"),
]
add_table(doc, ["Phương pháp", "Cấu hình chính", "Vai trò"], model_rows,
          widths=[1.70, 2.15, 2.95], font_size=9.2)

doc.add_heading("9 1 Cách chia dữ liệu", level=2)
add_paragraph(doc,
    "Dữ liệu phải được sắp theo thời gian và nhóm theo sự kiện mưa. Nhóm sử dụng các sự kiện đầu để huấn "
    "luyện các sự kiện tiếp theo để validation và giữ lại ít nhất một nhóm sự kiện cuối làm test. Nếu số "
    "sự kiện nguy hiểm quá ít kết quả phải được trình bày là thử nghiệm sơ bộ kèm số mẫu thực tế.")

doc.add_heading("9 2 Chỉ số đánh giá", level=2)
metric_rows = [
    ("Recall lớp 2", "Tỷ lệ sự kiện nguy hiểm được phát hiện", "Chỉ số an toàn chính"),
    ("Precision lớp 2", "Tỷ lệ cảnh báo đỏ là đúng", "Kiểm soát cảnh báo giả"),
    ("F1 Macro", "Cân bằng ba lớp", "Không để lớp an toàn lấn át"),
    ("Lead time", "Số phút cảnh báo trước khi vượt ngưỡng", "Giá trị vận hành"),
    ("Cảnh báo giả", "Số lần cảnh báo sai trên mỗi sự kiện mưa", "Khả năng sử dụng thực tế"),
    ("Độ sẵn sàng", "Tỷ lệ bản ghi hợp lệ và thời gian online", "Chất lượng hệ thống IoT"),
]
add_table(doc, ["Chỉ số", "Cách hiểu", "Mục đích"], metric_rows,
          widths=[1.40, 3.25, 2.15], font_size=9.0)

doc.add_heading("9 3 Tiêu chí thành công thử nghiệm", level=2)
for x in [
    "Pipeline ghi nhận được ít nhất 90 phần trăm chu kỳ dự kiến trong giai đoạn thử nghiệm.",
    "Mô hình vượt baseline trên Recall lớp nguy hiểm và F1 Macro của tập sự kiện độc lập.",
    "Thời gian cảnh báo trung vị đạt ít nhất 15 phút trên các sự kiện phát hiện đúng.",
    "Mỗi kết quả báo cáo kèm số mẫu số sự kiện và ma trận nhầm lẫn thay vì chỉ nêu Accuracy.",
]:
    add_bullet(doc, x)

doc.add_heading("10 Dashboard và cảnh báo", level=1)
for x in [
    "Hiển thị mưa mực nước tốc độ dâng và trạng thái kết nối theo cùng dấu thời gian.",
    "Tách rõ dự báo 15 phút dự báo 30 phút và cảnh báo bối cảnh 24 giờ.",
    "Mỗi cảnh báo phải ghi model version input timestamp xác suất và kết quả gửi Telegram.",
    "Cho phép người phụ trách xác nhận cảnh báo đúng sai để tạo dữ liệu phản hồi.",
    "Khi mất model hoặc mất Internet hệ thống chuyển sang luật ngưỡng an toàn và báo lỗi thiết bị.",
]:
    add_bullet(doc, x)

doc.add_heading("11 Kế hoạch triển khai", level=1)
schedule_rows = [
    ("Tuần 5", "Khảo sát điểm ngập và xác định khoảng đo", "Sơ đồ vị trí và biên bản ngưỡng"),
    ("Tuần 6", "Kiểm tra nguồn RS485 lắp cảm biến mưa và bổ sung cảm biến mực nước", "Mạch đo hoạt động an toàn trên bàn"),
    ("Tuần 7", "Hiệu chuẩn thiết bị xây HTTP hoặc MQTT ingestion", "Dữ liệu thật hiển thị trên dashboard"),
    ("Tuần 8", "Thu thập dữ liệu và tổ chức thử nghiệm máng nước", "Dataset có nguồn và cờ chất lượng"),
    ("Tuần 9", "Tạo nhãn huấn luyện baseline và ba mô hình", "Bảng benchmark theo thời gian"),
    ("Tuần 10", "Kiểm định sự kiện độc lập hoàn thiện báo cáo và diễn tập", "Prototype báo cáo poster và video minh chứng"),
]
add_table(doc, ["Thời gian", "Nhiệm vụ", "Sản phẩm nghiệm thu"], schedule_rows,
          widths=[1.05, 3.25, 2.50], font_size=9.0)

doc.add_heading("12 Rủi ro và biện pháp kiểm soát", level=1)
risk_rows = [
    ("Thiếu sự kiện ngập thật", "Không đủ lớp nguy hiểm", "Báo cáo nghiên cứu sơ bộ kéo dài thu thập không dùng synthetic làm test"),
    ("Cảm biến mưa sai thanh ghi", "Dữ liệu sai đơn vị", "Đọc datasheet và đối chứng thủ công"),
    ("Module TTL 5V", "Hỏng GPIO ESP32", "Dùng MAX3485 hoặc chuyển mức logic"),
    ("Mất WiFi hoặc nguồn", "Mất bản ghi và cảnh báo", "Bộ đệm cục bộ watchdog và báo trạng thái offline"),
    ("Nhiễu hoặc rác làm nghẽn", "Mực nước bất thường", "Median filter cờ chất lượng và vệ sinh định kỳ"),
    ("Cảnh báo quá mức", "Người dùng mất tin tưởng", "Đo Precision cooldown và xác nhận cảnh báo"),
]
add_table(doc, ["Rủi ro", "Ảnh hưởng", "Kiểm soát"], risk_rows,
          widths=[1.45, 1.80, 3.55], font_size=8.9)

doc.add_heading("13 Sản phẩm dự kiến", level=1)
for x in [
    "Một trạm ESP32 đọc cảm biến mưa RS485 và cảm biến mực nước.",
    "Một kho dữ liệu chuỗi thời gian có tài liệu mô tả trường và nguồn dữ liệu.",
    "Các mô hình baseline Logistic Regression Decision Tree và Random Forest đã đánh giá theo thời gian.",
    "Dashboard giám sát và cảnh báo Telegram có nhật ký kết quả gửi.",
    "Báo cáo khoa học poster sơ đồ mạch mã nguồn và video thử nghiệm.",
]:
    add_bullet(doc, x)

doc.add_heading("14 Kết luận", level=1)
add_paragraph(doc,
    "Phiên bản 2 chuyển đề tài từ một mô hình minh họa chung thành một nghiên cứu có thể kiểm chứng tại "
    "một điểm ngập cụ thể. Trọng tâm của đề tài là đo đúng gán nhãn đúng và đánh giá trên sự kiện chưa "
    "từng thấy. AI chỉ được xem là có giá trị khi vượt baseline và tạo được thời gian báo trước thực tế mà "
    "không làm số cảnh báo giả tăng đến mức không thể sử dụng.")

doc.add_heading("Tài liệu và dữ liệu nền", level=1)
refs = [
    "Bản đề xuất dự án nghiên cứu khoa học kỹ thuật phiên bản 1 năm học 2026 2027.",
    "Mã nguồn dataset hiện tại và kho bản tin khí tượng thủy văn đã lưu trong dự án.",
    "Dữ liệu API phải được ghi nguồn trong metadata và dữ liệu lưu lượng mô phỏng không được gọi là quan trắc thực địa.",
]
for ref in refs:
    add_bullet(doc, ref)

# Document properties
doc.core_properties.title = "Bản đề xuất cải tiến hệ thống dự báo ngập cục bộ tại trường học"
doc.core_properties.subject = "Đề xuất nghiên cứu khoa học kỹ thuật phiên bản 2"
doc.core_properties.author = "Nhóm dự án Trường THPT Trần Quý Cáp"
doc.core_properties.keywords = "ESP32, RS485, dự báo ngập, học máy, IoT"

doc.save(OUT_FILE)
print(OUT_FILE)
